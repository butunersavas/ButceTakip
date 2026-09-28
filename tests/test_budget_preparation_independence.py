import unittest
from decimal import Decimal
from io import BytesIO

from fastapi import UploadFile
from openpyxl import load_workbook
from sqlalchemy.pool import StaticPool
from sqlmodel import Session, SQLModel, create_engine, select

from app.models import BudgetItem, BudgetPreparation, PlanEntry, Scenario, User
from app.routers.budget_preparations import (
    create_item,
    create_preparation,
    export_preparation_xlsx,
    list_carryovers,
    mark_preparation_ready,
    reopen_preparation,
    update_item,
)
from app.routers.dashboard import get_dashboard
from app.schemas import BudgetPreparationCreate, BudgetPreparationItemInput
from app.services.budget_preparation_export import BUDGET_PREPARATION_EXPORT_HEADERS
from app.services.importer import import_xlsx


class BudgetPreparationIndependenceTests(unittest.TestCase):
    def setUp(self) -> None:
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        SQLModel.metadata.create_all(self.engine)
        self.session = Session(self.engine)
        self.user = User(username="prep-editor", hashed_password="test", role="user", is_active=True)
        self.session.add(self.user)
        self.session.commit()

    def tearDown(self) -> None:
        self.session.close()
        self.engine.dispose()

    def create_draft(self, year: int = 2027, name: str = "TEST01") -> BudgetPreparation:
        result = create_preparation(
            BudgetPreparationCreate(year=year, name=name, currency="USD"),
            self.session,
            self.user,
        )
        return self.session.get(BudgetPreparation, result.id)

    def cross_year_payload(self, *, description: str = "Yıl aşan test") -> BudgetPreparationItemInput:
        return BudgetPreparationItemInput(
            budget_name="TEST01",
            total_amount=Decimal("12565.00"),
            currency="USD",
            capex_opex="CAPEX",
            department="Sistem",
            map_attribute="Donanım",
            description=description,
            distribution_method="EQUAL",
            start_month=3,
            month_count=12,
            allocations=[],
        )

    def test_carryover_reads_preparation_allocations_not_operational_plans(self) -> None:
        draft = self.create_draft()
        create_item(draft.id, self.cross_year_payload(), self.session, self.user)
        scenario = Scenario(name="Operasyonel", year=2027, is_primary=True)
        budget_item = BudgetItem(code="OPER-1", name="Operasyonel Kalem")
        self.session.add(scenario)
        self.session.add(budget_item)
        self.session.flush()
        self.session.add(PlanEntry(
            year=2028,
            month=1,
            amount=999999,
            scenario_id=scenario.id,
            budget_item_id=budget_item.id,
        ))
        self.session.commit()

        draft_result = list_carryovers(2028, self.session, self.user)
        mark_preparation_ready(draft.id, self.session, self.user)
        ready_result = list_carryovers(2028, self.session, self.user)

        self.assertEqual(1, len(draft_result))
        self.assertEqual(1, len(ready_result))
        self.assertEqual(draft.id, ready_result[0].source_preparation_id)
        self.assertIsNone(ready_result[0].source_scenario_id)
        self.assertEqual([1, 2], [row.month for row in ready_result[0].months])
        self.assertEqual(Decimal("2094.20"), ready_result[0].total_amount)

    def test_export_contract_is_numeric_cross_year_and_import_compatible(self) -> None:
        draft = self.create_draft()
        create_item(draft.id, self.cross_year_payload(), self.session, self.user)
        response = export_preparation_xlsx(draft.id, self.session, self.user)
        worksheet = load_workbook(BytesIO(response.body), data_only=False)["Örnek"]

        self.assertEqual(BUDGET_PREPARATION_EXPORT_HEADERS, [cell.value for cell in worksheet[1]])
        self.assertEqual(13, worksheet.max_row)
        self.assertEqual({2027, 2028}, {worksheet.cell(row=row, column=5).value for row in range(2, 14)})
        for row in range(2, 14):
            for column in (5, 6, 7, 9, 10):
                self.assertEqual("n", worksheet.cell(row=row, column=column).data_type)
        self.assertEqual("Temel", worksheet["D2"].value)
        self.assertEqual("YANLIŞ", worksheet["N2"].value)

        summary = import_xlsx(
            UploadFile(filename="budget-preparation.xlsx", file=BytesIO(response.body)),
            self.session,
        )
        self.assertEqual(12, summary.imported_plans)

    def test_full_preparation_workflow_leaves_dashboard_plans_and_scenarios_unchanged(self) -> None:
        scenario = Scenario(name="Operasyonel Temel", year=2027, is_primary=True)
        budget_item = BudgetItem(code="BASE-1", name="Mevcut Plan")
        self.session.add(scenario)
        self.session.add(budget_item)
        self.session.flush()
        self.session.add(PlanEntry(
            year=2027,
            month=1,
            amount=4321.50,
            scenario_id=scenario.id,
            budget_item_id=budget_item.id,
        ))
        self.session.commit()
        dashboard_before = get_dashboard(
            year=2027,
            scenario_id=scenario.id,
            month=None,
            month_list=None,
            budget_item_id=None,
            department=None,
            capex_opex=None,
            session=self.session,
            _=self.user,
        )
        scenario_count = len(self.session.exec(select(Scenario)).all())
        plan_count = len(self.session.exec(select(PlanEntry)).all())

        draft = self.create_draft()
        item = create_item(draft.id, self.cross_year_payload(), self.session, self.user)
        mark_preparation_ready(draft.id, self.session, self.user)
        export_preparation_xlsx(draft.id, self.session, self.user)
        reopened = reopen_preparation(draft.id, self.session, self.user)
        update_item(
            draft.id,
            item.id,
            self.cross_year_payload(description="Düzenlemeye açıldıktan sonra güncellendi"),
            self.session,
            self.user,
        )

        dashboard_after = get_dashboard(
            year=2027,
            scenario_id=scenario.id,
            month=None,
            month_list=None,
            budget_item_id=None,
            department=None,
            capex_opex=None,
            session=self.session,
            _=self.user,
        )
        self.assertEqual(draft.id, reopened.id)
        self.assertEqual("DRAFT", reopened.status)
        self.assertEqual(scenario_count, len(self.session.exec(select(Scenario)).all()))
        self.assertEqual(plan_count, len(self.session.exec(select(PlanEntry)).all()))
        self.assertAlmostEqual(dashboard_before.kpi.total_plan, dashboard_after.kpi.total_plan, places=2)


if __name__ == "__main__":
    unittest.main()
