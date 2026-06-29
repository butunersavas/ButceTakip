from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import from_excel
from pydantic import ValidationError
from sqlmodel import Session, func, select

from app.models import User, WarrantyItem, WarrantyItemType
from app.schemas import WarrantyItemCreate

MODEL_FIELDS = (
    "type",
    "name",
    "location",
    "domain",
    "end_date",
    "note",
    "issuer",
    "certificate_issuer",
    "renewal_owner",
    "renewal_responsible",
    "purchased_from",
    "brand",
    "model",
    "serial_number",
    "asset_tag",
    "service_code",
    "ordered_product_model",
    "price",
    "shipment_date",
    "end_of_service_life",
    "status",
    "reminder_days",
    "remind_days",
    "remind_days_before",
)
REQUIRED_FIELDS = {"name"}
DATE_FIELDS = {"end_date", "shipment_date", "end_of_service_life"}
DECIMAL_FIELDS = {"price"}
INTEGER_FIELDS = {"reminder_days", "remind_days", "remind_days_before"}
TEXT_FIELDS = {
    "name",
    "location",
    "domain",
    "note",
    "issuer",
    "certificate_issuer",
    "renewal_owner",
    "renewal_responsible",
    "purchased_from",
    "brand",
    "model",
    "serial_number",
    "asset_tag",
    "service_code",
    "ordered_product_model",
    "status",
}

ALLOWED_FIELDS_BY_TYPE: dict[WarrantyItemType, set[str]] = {
    WarrantyItemType.DEVICE: {
        "type",
        "name",
        "purchased_from",
        "brand",
        "model",
        "serial_number",
        "asset_tag",
        "service_code",
        "ordered_product_model",
        "price",
        "shipment_date",
        "end_date",
        "end_of_service_life",
        "status",
        "note",
    },
    WarrantyItemType.DOMAIN_SSL: {
        "type",
        "name",
        "domain",
        "end_date",
        "renewal_responsible",
        "renewal_owner",
        "purchased_from",
    },
    WarrantyItemType.SERVICE: {
        "type",
        "name",
        "purchased_from",
        "service_code",
        "price",
        "shipment_date",
        "end_date",
        "status",
    },
}

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "type": ("type", "tip", "tur", "tür", "kayit_tipi", "kayıt_tipi", "kategori", "garanti_tipi"),
    "name": (
        "name",
        "ad",
        "adi",
        "adı",
        "baslik",
        "başlık",
        "urun",
        "ürün",
        "cihaz",
        "cihaz_adi",
        "cihaz_adı",
        "urun_adi",
        "ürün_adı",
        "hizmet",
        "hizmet_adi",
        "hizmet_adı",
        "garanti_adi",
        "bakim_adi",
    ),
    "location": ("location", "lokasyon", "konum", "yer", "departman", "bolum", "bölüm", "tesis", "adres"),
    "domain": ("domain", "domain_adi", "domain_adı", "domain_adlari", "domain_adları", "domainler", "alan_adi", "alan_adı", "fqdn", "url", "site"),
    "end_date": (
        "end_date",
        "enddate",
        "bitis_tarihi",
        "bitiş_tarihi",
        "garanti_bitis_tarihi",
        "bakim_bitis_tarihi",
        "son_tarih",
        "expiry_date",
        "expiration_date",
        "valid_until",
        "sozlesme_bitis_tarihi",
        "sözleşme_bitiş_tarihi",
        "sozlesme_sonu_tarihi",
        "destek_sonu_tarihi",
        "destek_sonu",
        "support_end_date",
    ),
    "note": ("note", "notes", "not", "aciklama", "açıklama", "description"),
    "issuer": ("issuer", "saglayici", "sağlayıcı", "duzenleyen", "düzenleyen", "veren"),
    "certificate_issuer": (
        "certificate_issuer",
        "certificateissuer",
        "sertifika_saglayici",
        "ssl_saglayici",
        "sertifika_veren",
    ),
    "renewal_owner": ("renewal_owner", "yenileme_sahibi", "owner", "kayit_sahibi", "kayıt_sahibi"),
    "renewal_responsible": (
        "renewal_responsible",
        "renewalresponsible",
        "yenileme_sorumlusu",
        "ilgili_firma",
        "ilgili_kurum",
        "sorumlu",
        "ilgili_kisi",
        "ilgili_kişi",
    ),
    "reminder_days": ("reminder_days", "hatirlatma_gun", "hatırlatma_gün", "hatirlatma", "kaç_gün_önce"),
    "remind_days": ("remind_days", "reminddays"),
    "remind_days_before": ("remind_days_before", "reminddaysbefore", "hatirlatma_gun_once"),
    "purchased_from": (
        "purchased_from",
        "purchasedfrom",
        "alinan_kurum",
        "alınan_kurum",
        "alinan_yer",
        "alınan_yer",
        "kurum",
        "vendor",
        "satici",
        "satıcı",
        "tedarikci",
        "tedarikçi",
        "supplier",
        "seller",
        "bayi",
        "hizmet_alinan_hosting_firmasi",
        "hizmet_alınan_hosting_firması",
        "hosting_firmasi",
        "hosting_firması",
    ),
    "brand": ("brand", "marka"),
    "model": ("model", "urun_modeli", "ürün_modeli", "cihaz_modeli"),
    "serial_number": (
        "serial_number",
        "serialnumber",
        "serial_no",
        "serial",
        "seri_no",
        "seri_numarasi",
        "seri_numarası",
    ),
    "asset_tag": ("asset_tag", "assettag", "demirbas", "demirbaş", "demirbas_no", "demirbaş_no", "envanter_no"),
    "service_code": ("service_code", "servicecode", "ekspres_servis_kodu", "express_service_code", "ekspres_kod", "lisans_adedi", "lisans_sayisi", "lisans_sayısı", "adet"),
    "ordered_product_model": ("ordered_product_model", "orderedproductmodel", "ordered_product"),
    "price": ("price", "fiyat", "amount", "tutar", "bedel", "maliyet", "ucret", "ücret", "cost"),
    "shipment_date": (
        "shipment_date",
        "shipmentdate",
        "gonderim_tarihi",
        "gönderim_tarihi",
        "alim_tarihi",
        "alım_tarihi",
        "sevkiyat_tarihi",
        "kargo_tarihi",
    ),
    "end_of_service_life": (
        "end_of_service_life",
        "endofservicelife",
        "eosl",
        "end_of_service",
        "servis_omru_sonu",
        "servis_ömrü_sonu",
    ),
    "status": ("status", "durum", "state", "garanti_suresi_uzatma_islemi_yapildi_mi", "garanti_süresi_uzatma_işlemi_yapıldı_mı", "uzatma_islemi", "uzatma_işlemi"),
}

TYPE_ALIASES = {
    "device": WarrantyItemType.DEVICE,
    "cihaz": WarrantyItemType.DEVICE,
    "urun": WarrantyItemType.DEVICE,
    "ürün": WarrantyItemType.DEVICE,
    "donanim": WarrantyItemType.DEVICE,
    "donanım": WarrantyItemType.DEVICE,
    "service": WarrantyItemType.SERVICE,
    "bakim": WarrantyItemType.SERVICE,
    "bakım": WarrantyItemType.SERVICE,
    "bakim_hizmet": WarrantyItemType.SERVICE,
    "hizmet": WarrantyItemType.SERVICE,
    "maintenance": WarrantyItemType.SERVICE,
    "lisans": WarrantyItemType.SERVICE,
    "lisans_destek": WarrantyItemType.SERVICE,
    "yazilim_lisans_destek": WarrantyItemType.SERVICE,
    "yazılım_lisans_destek": WarrantyItemType.SERVICE,
    "domain": WarrantyItemType.DOMAIN_SSL,
    "domain_ssl": WarrantyItemType.DOMAIN_SSL,
    "ssl": WarrantyItemType.DOMAIN_SSL,
    "sertifika": WarrantyItemType.DOMAIN_SSL,
}

TOTAL_MARKERS = {"toplam", "genel_toplam", "ara_toplam", "subtotal", "grand_total", "total"}
TURKISH_TRANSLATION = str.maketrans(
    {
        "İ": "I",
        "ı": "i",
        "Ş": "S",
        "ş": "s",
        "Ğ": "G",
        "ğ": "g",
        "Ü": "U",
        "ü": "u",
        "Ö": "O",
        "ö": "o",
        "Ç": "C",
        "ç": "c",
    }
)
TEMPLATE_COLUMNS_BY_TYPE = {
    WarrantyItemType.DEVICE: [
        "Alınan Kurum",
        "Ürün",
        "Marka",
        "Model",
        "Seri No",
        "Demirbaş",
        "Ekspres Servis Kodu",
        "Ordered Product Model",
        "Fiyat",
        "Gönderim Tarihi",
        "Destek Sonu Tarihi",
        "End of Service Life",
        "Durum",
        "Not",
    ],
    WarrantyItemType.DOMAIN_SSL: [
        "DOMAİN ADLARI",
        "SÖZLEŞME BİTİŞ TARİHİ",
        "SÖZLEŞME KALAN GÜN SAYISI",
        "İLGİLİ FİRMA",
        "HİZMET ALINAN HOSTİNG FİRMASI",
    ],
    WarrantyItemType.SERVICE: [
        "Alınan Kurum",
        "Ürün",
        "Lisans Adedi",
        "Fiyat",
        "Alım Tarihi",
        "Bitiş Tarihi",
        "Destek Kalan Gün",
        "Garanti Süresi Uzatma İşlemi Yapıldı Mı?",
    ],
}

TEMPLATE_COLUMNS = TEMPLATE_COLUMNS_BY_TYPE[WarrantyItemType.DEVICE]


@dataclass
class RowResult:
    row_number: int
    status: str
    payload: dict[str, Any] = field(default_factory=dict)
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    duplicate_key: tuple[str, ...] | None = None
    duplicate_source: str | None = None
    amount: Decimal | None = None
    shipment_date: date | None = None
    end_date: date | None = None
    end_of_service_life: date | None = None


def json_safe(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, WarrantyItemType):
        return value.value
    if isinstance(value, tuple):
        return [json_safe(part) for part in value]
    if isinstance(value, list):
        return [json_safe(part) for part in value]
    if isinstance(value, dict):
        return {key: json_safe(part) for key, part in value.items()}
    return value


def normalize_key(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().translate(TURKISH_TRANSLATION)
    text = unicodedata.normalize("NFKD", text)
    text = text.encode("ascii", "ignore").decode("ascii")
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = value.strip() if isinstance(value, str) else str(value).strip()
    if not text or normalize_key(text) in {"", "nan", "none", "null"}:
        return None
    if text in {"-", "--", "—"}:
        return None
    return text


def build_header_lookup() -> tuple[dict[str, str], dict[str, set[str]]]:
    lookup: dict[str, str] = {}
    conflicts: dict[str, set[str]] = {}
    for field_name, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            key = normalize_key(alias)
            existing = lookup.get(key)
            if existing and existing != field_name:
                conflicts.setdefault(key, {existing}).add(field_name)
            else:
                lookup[key] = field_name
    for conflict_key in conflicts:
        lookup.pop(conflict_key, None)
    return lookup, conflicts


HEADER_LOOKUP, HEADER_CONFLICTS = build_header_lookup()


def read_workbook_from_bytes(content: bytes, filename: str, sheet_name: str | None = None) -> tuple[str, list[list[Any]]]:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Sadece .xlsx veya .xlsm dosyaları destekleniyor. .xls dosyasını .xlsx olarak kaydedip tekrar deneyin.")
    workbook = load_workbook(BytesIO(content), data_only=True, read_only=True)
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"'{sheet_name}' sayfası bulunamadı. Mevcut sayfalar: {', '.join(workbook.sheetnames)}")
            worksheet = workbook[sheet_name]
        else:
            worksheet = workbook["Garanti Bakım Import"] if "Garanti Bakım Import" in workbook.sheetnames else workbook[workbook.sheetnames[0]]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        return worksheet.title, rows
    finally:
        workbook.close()


def get_cell(row: list[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def is_empty_row(row: list[Any]) -> bool:
    return all(clean_text(value) is None for value in row)


def is_total_row(row: list[Any]) -> bool:
    non_empty = [clean_text(value) for value in row if clean_text(value) is not None]
    if not non_empty:
        return False
    first_key = normalize_key(non_empty[0])
    if first_key in TOTAL_MARKERS or "toplam" in first_key:
        return True
    if len(non_empty) <= 2:
        return any(normalize_key(value) in TOTAL_MARKERS for value in non_empty)
    return False


def find_header_row(rows: list[list[Any]]) -> tuple[int, list[Any]]:
    for index, row in enumerate(rows[:25]):
        normalized = [normalize_key(value) for value in row if clean_text(value)]
        mapped = [key for key in normalized if key in HEADER_LOOKUP or key in HEADER_CONFLICTS or key in TOTAL_MARKERS]
        if len(mapped) >= 2 and not is_total_row(row):
            return index, row
    raise ValueError("İlk 25 satır içinde kolon başlık satırı bulunamadı.")


def map_headers(headers: list[Any]) -> dict[str, Any]:
    mapped: dict[str, int] = {}
    mapped_headers: dict[str, str] = {}
    unmapped: list[str] = []
    ambiguous: list[dict[str, Any]] = []
    duplicate_mappings: list[dict[str, Any]] = []

    for index, header in enumerate(headers):
        header_text = clean_text(header)
        if not header_text:
            continue
        key = normalize_key(header_text)
        if key in HEADER_CONFLICTS:
            ambiguous.append({"column": header_text, "normalized": key, "candidate_fields": sorted(HEADER_CONFLICTS[key])})
            continue
        field_name = HEADER_LOOKUP.get(key)
        if not field_name:
            unmapped.append(header_text)
            continue
        if field_name in mapped:
            duplicate_mappings.append({"field": field_name, "kept_column": mapped_headers[field_name], "ignored_column": header_text})
            continue
        mapped[field_name] = index
        mapped_headers[field_name] = header_text
    return {
        "indexes": mapped,
        "headers": mapped_headers,
        "unmapped_columns": unmapped,
        "ambiguous_columns": ambiguous,
        "duplicate_mappings": duplicate_mappings,
    }


def parse_type(value: Any, default_type: WarrantyItemType | None = WarrantyItemType.DEVICE) -> WarrantyItemType | None:
    text = clean_text(value)
    if not text:
        return default_type
    key = normalize_key(text)
    if key in TYPE_ALIASES:
        return TYPE_ALIASES[key]
    try:
        return WarrantyItemType(text)
    except ValueError:
        try:
            return WarrantyItemType(text.upper())
        except ValueError:
            return None


def parse_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        try:
            return from_excel(value).date()
        except Exception:
            return None
    text = clean_text(value)
    if not text:
        return None
    text = text.replace("\\", "/")
    iso_candidate = text[:10]
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%m/%d/%Y"):
        try:
            candidate = iso_candidate if fmt == "%Y-%m-%d" else text
            return datetime.strptime(candidate, fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        return None


def parse_decimal(value: Any) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = clean_text(value)
    if not text:
        return None
    text = text.replace("\u00a0", " ")
    negative = text.startswith("(") and text.endswith(")")
    cleaned = re.sub(r"(?i)(try|tl|usd|eur|gbp|dolar|euro)", "", text)
    cleaned = re.sub(r"[^0-9,.\-]", "", cleaned)
    if not cleaned or cleaned in {"-", ".", ","}:
        return None
    if "," in cleaned and "." in cleaned:
        decimal_separator = "," if cleaned.rfind(",") > cleaned.rfind(".") else "."
        thousands_separator = "." if decimal_separator == "," else ","
        cleaned = cleaned.replace(thousands_separator, "").replace(decimal_separator, ".")
    elif "," in cleaned:
        decimals = cleaned.rsplit(",", 1)[1]
        cleaned = cleaned.replace(".", "")
        cleaned = cleaned.replace(",", "." if len(decimals) in {1, 2} else "")
    elif "." in cleaned:
        decimals = cleaned.rsplit(".", 1)[1]
        if len(decimals) not in {1, 2}:
            cleaned = cleaned.replace(".", "")
    try:
        amount = Decimal(cleaned)
    except InvalidOperation:
        return None
    return -amount if negative else amount


def parse_int(value: Any) -> int | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    text = clean_text(value)
    if not text:
        return None
    try:
        return int(Decimal(text.replace(",", ".")))
    except InvalidOperation:
        return None


def normalize_warranty_payload(payload: dict[str, Any]) -> dict[str, Any]:
    if payload.get("certificate_issuer") and not payload.get("issuer"):
        payload["issuer"] = payload["certificate_issuer"]
    if payload.get("issuer") and not payload.get("certificate_issuer"):
        payload["certificate_issuer"] = payload["issuer"]
    if payload.get("renewal_responsible") and not payload.get("renewal_owner"):
        payload["renewal_owner"] = payload["renewal_responsible"]
    if payload.get("renewal_owner") and not payload.get("renewal_responsible"):
        payload["renewal_responsible"] = payload["renewal_owner"]
    reminder_value = payload.get("reminder_days")
    if reminder_value is not None:
        payload.setdefault("remind_days", reminder_value)
        payload.setdefault("remind_days_before", reminder_value)
    remind_value = payload.get("remind_days") or payload.get("remind_days_before")
    if remind_value is not None:
        payload.setdefault("reminder_days", remind_value)
        payload.setdefault("remind_days", remind_value)
        payload.setdefault("remind_days_before", remind_value)
    return payload


def restrict_payload_for_type(payload: dict[str, Any], default_type: WarrantyItemType) -> dict[str, Any]:
    allowed_fields = ALLOWED_FIELDS_BY_TYPE.get(default_type, ALLOWED_FIELDS_BY_TYPE[WarrantyItemType.DEVICE])
    restricted = {key: value for key, value in payload.items() if key in allowed_fields}
    restricted["type"] = default_type
    return restricted


def stringify_key_part(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, WarrantyItemType):
        return value.value
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return normalize_key(value)


def make_duplicate_key(payload: dict[str, Any]) -> tuple[str, ...]:
    item_type = payload.get("type")
    if item_type == WarrantyItemType.DOMAIN_SSL or item_type == WarrantyItemType.DOMAIN_SSL.value:
        return (
            "domain",
            stringify_key_part(payload.get("domain") or payload.get("name")),
            stringify_key_part(payload.get("end_date")),
            stringify_key_part(payload.get("renewal_responsible") or payload.get("renewal_owner")),
            stringify_key_part(payload.get("purchased_from")),
        )
    if item_type == WarrantyItemType.SERVICE or item_type == WarrantyItemType.SERVICE.value:
        return (
            "license_support",
            stringify_key_part(payload.get("name")),
            stringify_key_part(payload.get("purchased_from")),
            stringify_key_part(payload.get("shipment_date")),
            stringify_key_part(payload.get("end_date")),
            stringify_key_part(payload.get("price")),
        )

    serial_number = stringify_key_part(payload.get("serial_number"))
    if serial_number:
        return ("device_serial_number", serial_number)
    asset_tag = stringify_key_part(payload.get("asset_tag"))
    if asset_tag:
        return ("device_asset_tag", asset_tag)
    service_code = stringify_key_part(payload.get("service_code"))
    if service_code:
        return ("device_service_code", service_code)
    return (
        "device_fallback",
        stringify_key_part(payload.get("name")),
        stringify_key_part(payload.get("purchased_from")),
        stringify_key_part(payload.get("end_date")),
        stringify_key_part(payload.get("price")),
    )


def warranty_to_payload(item: WarrantyItem) -> dict[str, Any]:
    return {
        "type": item.type,
        "name": item.name,
        "location": item.location,
        "domain": item.domain,
        "end_date": item.end_date,
        "issuer": item.issuer or item.certificate_issuer,
        "purchased_from": item.purchased_from,
        "renewal_responsible": item.renewal_responsible or item.renewal_owner,
        "renewal_owner": item.renewal_owner or item.renewal_responsible,
        "shipment_date": item.shipment_date,
        "serial_number": item.serial_number,
        "asset_tag": item.asset_tag,
        "service_code": item.service_code,
        "price": item.price,
    }


def load_existing_duplicate_keys(session: Session) -> dict[tuple[str, ...], int | None]:
    items = session.exec(select(WarrantyItem).where(WarrantyItem.is_active.is_(True))).all()
    return {make_duplicate_key(warranty_to_payload(item)): item.id for item in items}


def parse_row(row_number: int, row: list[Any], header_map: dict[str, int], default_type: WarrantyItemType) -> RowResult:
    payload: dict[str, Any] = {}
    errors: list[str] = []
    warnings: list[str] = []
    amount: Decimal | None = None
    shipment_date: date | None = None
    end_date: date | None = None
    end_of_service_life: date | None = None

    # Import ekranında seçilen sekme kaynak tiptir. Excel içindeki Tip kolonu varsa bile
    # başka kayıt türüne yazılmasın; Cihaz/Domain/Lisans Destek birbirine karışmasın.
    payload["type"] = default_type

    for field_name in MODEL_FIELDS:
        if field_name == "type":
            continue
        cell_value = get_cell(row, header_map.get(field_name))
        if field_name in DATE_FIELDS:
            parsed_date = parse_date(cell_value)
            if parsed_date:
                payload[field_name] = parsed_date
                if field_name == "end_date":
                    end_date = parsed_date
                elif field_name == "shipment_date":
                    shipment_date = parsed_date
                elif field_name == "end_of_service_life":
                    end_of_service_life = parsed_date
            elif clean_text(cell_value) is not None:
                errors.append(f"{field_name} tarihi okunamadı.")
        elif field_name in DECIMAL_FIELDS:
            parsed_decimal = parse_decimal(cell_value)
            if parsed_decimal is not None:
                payload[field_name] = parsed_decimal
                if field_name == "price":
                    amount = parsed_decimal
            elif clean_text(cell_value) is not None:
                errors.append(f"{field_name} Decimal/sayı olarak okunamadı.")
        elif field_name in INTEGER_FIELDS:
            parsed_int = parse_int(cell_value)
            if parsed_int is not None:
                if parsed_int < 0:
                    errors.append(f"{field_name} negatif olamaz.")
                else:
                    payload[field_name] = parsed_int
        elif field_name in TEXT_FIELDS:
            parsed_text = clean_text(cell_value)
            if parsed_text is not None:
                payload[field_name] = parsed_text

    missing_values = [field_name for field_name in REQUIRED_FIELDS if payload.get(field_name) in {None, ""}]
    if missing_values:
        errors.append("Zorunlu alan eksik: " + ", ".join(sorted(missing_values)))

    payload = normalize_warranty_payload(payload)
    if default_type == WarrantyItemType.DOMAIN_SSL:
        if not payload.get("name") and payload.get("domain"):
            payload["name"] = payload["domain"]
        if not payload.get("domain") and payload.get("name"):
            payload["domain"] = payload["name"]
    if default_type == WarrantyItemType.SERVICE and payload.get("service_code") is not None:
        payload["service_code"] = str(payload["service_code"])
    payload = restrict_payload_for_type(payload, default_type)
    if default_type == WarrantyItemType.DOMAIN_SSL:
        if not payload.get("name") and payload.get("domain"):
            payload["name"] = payload["domain"]
        if not payload.get("domain") and payload.get("name"):
            payload["domain"] = payload["name"]
    if not errors:
        try:
            validated = WarrantyItemCreate(**payload)
            payload = validated.dict()
        except ValidationError as exc:
            errors.extend(error["msg"] for error in exc.errors())

    return RowResult(
        row_number=row_number,
        status="invalid" if errors else "importable",
        payload=payload,
        errors=errors,
        warnings=warnings,
        amount=amount,
        shipment_date=shipment_date,
        end_date=end_date or payload.get("end_date"),
        end_of_service_life=end_of_service_life or payload.get("end_of_service_life"),
    )


def analyze_rows(
    rows: list[list[Any]],
    header_row_index: int,
    header_map: dict[str, int],
    default_type: WarrantyItemType,
    existing_keys: dict[tuple[str, ...], int | None],
) -> tuple[list[RowResult], dict[str, int], Decimal | None, dict[str, str | None]]:
    results: list[RowResult] = []
    seen_keys: dict[tuple[str, ...], int] = {}
    counts = {
        "worksheet_rows_after_header": max(len(rows) - header_row_index - 1, 0),
        "blank_rows": 0,
        "total_rows": 0,
        "data_rows_read": 0,
        "importable_rows": 0,
        "invalid_rows": 0,
        "duplicate_rows": 0,
    }
    total_amount: Decimal | None = None
    shipment_dates: list[date] = []
    end_dates: list[date] = []
    eol_dates: list[date] = []

    for row_index, row in enumerate(rows[header_row_index + 1 :], start=header_row_index + 2):
        if is_empty_row(row):
            counts["blank_rows"] += 1
            continue
        if is_total_row(row):
            counts["total_rows"] += 1
            continue

        counts["data_rows_read"] += 1
        result = parse_row(row_index, row, header_map, default_type)
        if result.status == "importable":
            duplicate_key = make_duplicate_key(result.payload)
            result.duplicate_key = duplicate_key
            if duplicate_key in existing_keys:
                result.status = "duplicate"
                result.duplicate_source = "database"
            elif duplicate_key in seen_keys:
                result.status = "duplicate"
                result.duplicate_source = f"file row {seen_keys[duplicate_key]}"
            else:
                seen_keys[duplicate_key] = row_index

        if result.status == "importable":
            counts["importable_rows"] += 1
            if result.amount is not None:
                total_amount = (total_amount or Decimal("0")) + result.amount
            if result.shipment_date:
                shipment_dates.append(result.shipment_date)
            if result.end_date:
                end_dates.append(result.end_date)
            if result.end_of_service_life:
                eol_dates.append(result.end_of_service_life)
        elif result.status == "duplicate":
            counts["duplicate_rows"] += 1
        else:
            counts["invalid_rows"] += 1
        results.append(result)

    date_range = {
        "shipment_date_min": min(shipment_dates).isoformat() if shipment_dates else None,
        "shipment_date_max": max(shipment_dates).isoformat() if shipment_dates else None,
        "end_date_min": min(end_dates).isoformat() if end_dates else None,
        "end_date_max": max(end_dates).isoformat() if end_dates else None,
        "end_of_service_life_min": min(eol_dates).isoformat() if eol_dates else None,
        "end_of_service_life_max": max(eol_dates).isoformat() if eol_dates else None,
    }
    return results, counts, total_amount, date_range


def result_preview(results: list[RowResult], status: str, limit: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for result in results:
        if result.status != status:
            continue
        rows.append(
            {
                "row": result.row_number,
                "status": result.status,
                "errors": result.errors,
                "warnings": result.warnings,
                "duplicate_source": result.duplicate_source,
                "duplicate_key": result.duplicate_key,
                "payload": result.payload,
            }
        )
        if len(rows) >= limit:
            break
    return rows


def get_active_warranty_count(session: Session) -> int:
    value = session.exec(select(func.count(WarrantyItem.id)).where(WarrantyItem.is_active.is_(True))).one()
    return int(value or 0)


def resolve_user_id(session: Session, user: User | None) -> int | None:
    if user is not None:
        return user.id
    admin_user = session.exec(select(User).where(User.is_active.is_(True), User.is_admin.is_(True))).first()
    return admin_user.id if admin_user else None


def build_report(
    *,
    mode: str,
    filename: str,
    sheet_name: str,
    header_row_index: int,
    headers: list[Any],
    mapping_info: dict[str, Any],
    missing_required_columns: list[str],
    results: list[RowResult],
    counts: dict[str, int],
    total_amount: Decimal | None,
    date_range: dict[str, str | None],
    db_total_before: int,
    db_total_after: int | None = None,
    added_rows: int = 0,
) -> dict[str, Any]:
    return json_safe(
        {
            "mode": mode,
            "file": filename,
            "sheet": sheet_name,
            "excel_columns": [clean_text(header) for header in headers if clean_text(header)],
            "header_row": header_row_index + 1,
            "model_fields": list(MODEL_FIELDS),
            "column_mapping": {field_name: mapping_info["headers"][field_name] for field_name in MODEL_FIELDS if field_name in mapping_info["headers"]},
            "unmapped_columns": mapping_info["unmapped_columns"],
            "ambiguous_columns": mapping_info["ambiguous_columns"],
            "duplicate_header_mappings": mapping_info["duplicate_mappings"],
            "missing_required_columns": missing_required_columns,
            "duplicate_control": {
                "implemented_key": [
                    "Cihaz: serial_number / asset_tag / ekspres_servis_kodu / fallback",
                    "Domain: domain + sözleşme_bitiş_tarihi + ilgili_firma + hosting_firması",
                    "Lisans Destek: ürün + alınan_kurum + alım_tarihi + bitiş_tarihi + fiyat",
                ],
                "note": "Seçilen sekmeye göre ayrı duplicate anahtarı kullanılır; Cihaz/Domain/Lisans Destek birbirine karışmaz.",
            },
            "summary": {
                **counts,
                "total_amount": total_amount,
                "date_range": date_range,
                "db_active_records_before": db_total_before,
                "db_active_records_after": db_total_after if db_total_after is not None else db_total_before,
                "added_rows": added_rows,
            },
            "invalid_row_preview": result_preview(results, "invalid", 50),
            "duplicate_row_preview": result_preview(results, "duplicate", 50),
            "importable_row_preview": result_preview(results, "importable", 25),
        }
    )


def preview_warranty_import(
    *,
    session: Session,
    content: bytes,
    filename: str,
    sheet_name: str | None = None,
    default_type: WarrantyItemType = WarrantyItemType.DEVICE,
) -> dict[str, Any]:
    sheet_name_resolved, rows = read_workbook_from_bytes(content, filename, sheet_name)
    header_row_index, headers = find_header_row(rows)
    mapping_info = map_headers(headers)
    missing_required_columns = [
        field_name for field_name in sorted(REQUIRED_FIELDS) if field_name not in mapping_info["headers"]
    ]
    if default_type == WarrantyItemType.DOMAIN_SSL and "name" in missing_required_columns and "domain" in mapping_info["headers"]:
        missing_required_columns.remove("name")
    db_total_before = get_active_warranty_count(session)

    if missing_required_columns:
        counts = {
            "worksheet_rows_after_header": max(len(rows) - header_row_index - 1, 0),
            "blank_rows": 0,
            "total_rows": 0,
            "data_rows_read": 0,
            "importable_rows": 0,
            "invalid_rows": 0,
            "duplicate_rows": 0,
        }
        return build_report(
            mode="dry_run",
            filename=filename,
            sheet_name=sheet_name_resolved,
            header_row_index=header_row_index,
            headers=headers,
            mapping_info=mapping_info,
            missing_required_columns=missing_required_columns,
            results=[],
            counts=counts,
            total_amount=None,
            date_range={
                "shipment_date_min": None,
                "shipment_date_max": None,
                "end_date_min": None,
                "end_date_max": None,
                "end_of_service_life_min": None,
                "end_of_service_life_max": None,
            },
            db_total_before=db_total_before,
        )

    existing_keys = load_existing_duplicate_keys(session)
    results, counts, total_amount, date_range = analyze_rows(
        rows,
        header_row_index,
        mapping_info["indexes"],
        default_type,
        existing_keys,
    )
    return build_report(
        mode="dry_run",
        filename=filename,
        sheet_name=sheet_name_resolved,
        header_row_index=header_row_index,
        headers=headers,
        mapping_info=mapping_info,
        missing_required_columns=missing_required_columns,
        results=results,
        counts=counts,
        total_amount=total_amount,
        date_range=date_range,
        db_total_before=db_total_before,
    )


def confirm_warranty_import(
    *,
    session: Session,
    content: bytes,
    filename: str,
    current_user: User | None = None,
    sheet_name: str | None = None,
    default_type: WarrantyItemType = WarrantyItemType.DEVICE,
) -> dict[str, Any]:
    sheet_name_resolved, rows = read_workbook_from_bytes(content, filename, sheet_name)
    header_row_index, headers = find_header_row(rows)
    mapping_info = map_headers(headers)
    missing_required_columns = [
        field_name for field_name in sorted(REQUIRED_FIELDS) if field_name not in mapping_info["headers"]
    ]
    if default_type == WarrantyItemType.DOMAIN_SSL and "name" in missing_required_columns and "domain" in mapping_info["headers"]:
        missing_required_columns.remove("name")
    db_total_before = get_active_warranty_count(session)
    existing_keys = load_existing_duplicate_keys(session)

    if missing_required_columns:
        return preview_warranty_import(session=session, content=content, filename=filename, sheet_name=sheet_name, default_type=default_type)

    results, counts, total_amount, date_range = analyze_rows(
        rows,
        header_row_index,
        mapping_info["indexes"],
        default_type,
        existing_keys,
    )

    actor_id = resolve_user_id(session, current_user)
    added_rows = 0
    try:
        for result in results:
            if result.status != "importable":
                continue
            payload = dict(result.payload)
            item = WarrantyItem(
                **payload,
                created_by_id=actor_id,
                updated_by_id=actor_id,
                created_by_user_id=actor_id,
                updated_by_user_id=actor_id,
            )
            session.add(item)
            added_rows += 1
        session.flush()
        db_total_after = get_active_warranty_count(session)
        session.commit()
    except Exception:
        session.rollback()
        raise

    return build_report(
        mode="commit",
        filename=filename,
        sheet_name=sheet_name_resolved,
        header_row_index=header_row_index,
        headers=headers,
        mapping_info=mapping_info,
        missing_required_columns=missing_required_columns,
        results=results,
        counts=counts,
        total_amount=total_amount,
        date_range=date_range,
        db_total_before=db_total_before,
        db_total_after=db_total_after,
        added_rows=added_rows,
    )


def build_template_workbook(template_type: WarrantyItemType = WarrantyItemType.DEVICE) -> BytesIO:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Garanti Bakım Import"
    columns = TEMPLATE_COLUMNS_BY_TYPE.get(template_type, TEMPLATE_COLUMNS)
    worksheet.append(columns)
    if template_type == WarrantyItemType.DOMAIN_SSL:
        worksheet.append([
            "ornekdomain.com",
            "31.12.2026",
            "",
            "İlgili Firma",
            "Hosting Firması",
        ])
    elif template_type == WarrantyItemType.SERVICE:
        worksheet.append([
            "Örnek Kurum",
            "Örnek Yazılım",
            "10",
            "1000,00",
            "01.01.2026",
            "31.12.2026",
            "",
            "Hayır",
        ])
    else:
        worksheet.append([
            "Örnek Kurum",
            "Örnek Ürün",
            "Marka",
            "Model",
            "SN123456",
            "DEM-001",
            "SRV-001",
            "Ordered Model",
            "1000,00",
            "01.01.2026",
            "31.12.2026",
            "31.12.2027",
            "Aktif",
            "Örnek satırdır, import öncesi silebilirsiniz.",
        ])
    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(max(max_length + 3, 14), 32)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
